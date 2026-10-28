#!/usr/bin/env python3
"""Presentación de la presustentación, con la plantilla de la presentación del anteproyecto.

Arma Documentos/Presentacion/Presustentacion.pptx con python-pptx. La plantilla (fondos, logos,
banda de acreditación e íconos) se extrajo del PDF de la presentación del anteproyecto y está en
Documentos/Presentacion/plantilla/, junto con cuadros del video de la campaña del 16 de octubre y
los códigos QR de los anexos. Las figuras son las del documento final
(Documentos/DocumentoFinal/imagenes/).

Las coordenadas van en puntos sobre la diapositiva de 960 × 540 pt de la plantilla, las mismas que
da pdftohtml sobre el PDF original. Cada diapositiva lleva en las notas el tiempo sugerido y lo que
se dice; el total ronda los 17 minutos.

Uso, desde la raíz del repositorio:
    python3 herramientas/presentacion_sustentacion.py
"""
import math
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

RAIZ = Path(__file__).resolve().parent.parent
REC = RAIZ / 'Documentos/Presentacion/plantilla'
FIG = RAIZ / 'Documentos/DocumentoFinal/imagenes'
SALIDA = RAIZ / 'Documentos/Presentacion/Presustentacion.pptx'

# Colores medidos sobre el PDF de la plantilla
AZUL_TITULO = '001F5F'
AZUL = '4472C4'
AZUL_CLARO = '85A3D8'
AZUL_ALCANCE = '5170FF'
TINTE = 'E7EDF7'
NARANJA_BORDE = 'D7712B'
DURAZNO = 'F8D7CD'
NARANJA = 'F6B79D'
AMARILLO = 'FFFF00'
VERDE = '00B050'
NEGRO = '000000'
GRIS = '404040'
GRIS_CLARO = 'D9D9D9'
BLANCO = 'FFFFFF'

ALIN = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT, 'j': PP_ALIGN.JUSTIFY}
ANCLA = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}

REPOSITORIO = ('https://github.com/shdeza12/Collaborative-Mobile-Robot-System-for-Indoor-'
               'Navigation-Assistance-in-Multi-Level-Environments')
SITIO = ('https://shdeza12.github.io/Collaborative-Mobile-Robot-System-for-Indoor-Navigation-'
         'Assistance-in-Multi-Level-Environments/')
VIDEOS = [  # anexo A del documento final
    ('Q0w_CTga98k', 'Seguidor de paredes en el pasillo simulado', '5 may.'),
    ('ZPJ0Qw2L5-A', 'Teleoperación y mapeo en el modelo de Gregorio', '19 may.'),
    ('ZGfAMnC4lYY', 'Prueba de marcha del AWS DeepRacer', '5 ago.'),
    ('6SABSgCxVEw', 'Media vuelta con Nav2 en simulación', '18 ago.'),
    ('iKYSiM9cs-Y', 'Primeras asignaciones del coordinador', '24 ago.'),
    ('eH_Co4vZZdw', 'Campaña de 30 misiones en simulación', '5 sep.'),
    ('X3BINGgNxiA', 'Campaña con los agentes físicos', '20 oct.'),
]


def rgb(h):
    return RGBColor.from_string(h)


# --------------------------------------------------------------- texto ---
def _vineta(p, sangria):
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', str(int(Pt(sangria))))
    pPr.set('indent', str(-int(Pt(sangria))))
    etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
    etree.SubElement(pPr, qn('a:buChar')).set('char', '•')


def escribir(tf, parrafos, tam=14, color=NEGRO, fuente='Arial', negrita=False, cursiva=False,
             alin='l', interlineado=None, despues=0, vineta=False, sangria=11):
    """Cada párrafo es un texto, o un dict con 't' (o 'runs': [(texto, opciones)]) y opciones."""
    if isinstance(parrafos, (str, dict)):
        parrafos = [parrafos]
    for i, par in enumerate(parrafos):
        if isinstance(par, str):
            par = {'t': par}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIN[par.get('alin', alin)]
        if par.get('interlineado', interlineado):
            p.line_spacing = par.get('interlineado', interlineado)
        if par.get('despues', despues):
            p.space_after = Pt(par.get('despues', despues))
        if par.get('vineta', vineta):
            _vineta(p, par.get('sangria', sangria))
        for seg in par.get('runs') or [(par['t'], {})]:
            t, op = (seg, {}) if isinstance(seg, str) else seg
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = op.get('fuente', par.get('fuente', fuente))
            f.size = Pt(op.get('tam', par.get('tam', tam)))
            f.bold = op.get('negrita', par.get('negrita', negrita))
            f.italic = op.get('cursiva', par.get('cursiva', cursiva))
            f.color.rgb = rgb(op.get('color', par.get('color', color)))
    return tf


def _marco(tf, anclaje, margen):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    m = margen if isinstance(margen, tuple) else (margen,) * 4
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (Pt(v) for v in m)
    tf.vertical_anchor = ANCLA[anclaje]


def texto(s, x, y, w, h, parrafos, anclaje='t', margen=0, **kw):
    caja = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    _marco(caja.text_frame, anclaje, margen)
    escribir(caja.text_frame, parrafos, **kw)
    return caja


def fuente(s, x, y, w, t='Fuente: Elaboración propia', alin='l'):
    return texto(s, x, y, w, 14, t, tam=9, color=GRIS, alin=alin)


# --------------------------------------------------------------- formas ---
def forma(s, tipo, x, y, w, h, relleno=None, borde=None, grosor=0.75, radio=None, guion=False,
          parrafos=None, anclaje='m', margen=4, **kw):
    sh = s.shapes.add_shape(tipo, Pt(x), Pt(y), Pt(w), Pt(h))
    estilo = sh._element.find(qn('p:style'))
    if estilo is not None:  # sin el estilo del tema: ni sombra ni texto blanco por omisión
        sh._element.remove(estilo)
    if relleno:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(relleno)
    else:
        sh.fill.background()
    if borde:
        sh.line.color.rgb = rgb(borde)
        sh.line.width = Pt(grosor)
        if guion:
            sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        sh.line.fill.background()
    if radio is not None:
        sh.adjustments[0] = radio
    if parrafos is not None:
        _marco(sh.text_frame, anclaje, margen)
        escribir(sh.text_frame, parrafos, **kw)
    return sh


def linea(s, puntos, color=AZUL, grosor=0.75):
    fb = s.shapes.build_freeform(puntos[0][0], puntos[0][1], scale=12700)
    fb.add_line_segments(puntos[1:], close=False)
    sh = fb.convert_to_shape()
    estilo = sh._element.find(qn('p:style'))
    if estilo is not None:
        sh._element.remove(estilo)
    sh.fill.background()
    sh.line.color.rgb = rgb(color)
    sh.line.width = Pt(grosor)
    return sh


def imagen(s, ruta, x, y, w=None, h=None, enlace=None):
    """Coloca la imagen dentro de la caja (w, h) sin deformarla, centrada; devuelve su caja."""
    iw, ih = Image.open(ruta).size
    if w and h:
        e = min(w / iw, h / ih)
        ww, hh = iw * e, ih * e
        x, y = x + (w - ww) / 2, y + (h - hh) / 2
    elif w:
        ww, hh = w, ih * w / iw
    else:
        ww, hh = iw * h / ih, h
    pic = s.shapes.add_picture(str(ruta), Pt(x), Pt(y), Pt(ww), Pt(hh))
    if enlace:
        pic.click_action.hyperlink.address = enlace
    return x, y, ww, hh


def tabla(s, x, y, anchos, filas, alto=24, tam=11, encabezado=True, altos=None):
    gf = s.shapes.add_table(len(filas), len(anchos), Pt(x), Pt(y), Pt(sum(anchos)),
                            Pt(alto * len(filas)))
    t = gf.table
    t.first_row = encabezado
    t.horz_banding = False
    for j, a in enumerate(anchos):
        t.columns[j].width = Pt(a)
    for i, fila in enumerate(filas):
        t.rows[i].height = Pt((altos or {}).get(i, alto))
        for j, valor in enumerate(fila):
            if valor is None:  # celda absorbida por la anterior
                continue
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Pt(5)
            c.margin_top = c.margin_bottom = Pt(2)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            cab = encabezado and i == 0
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(AZUL if cab else (BLANCO if i % 2 else TINTE))
            tf = c.text_frame
            tf.word_wrap = True
            pars = valor if isinstance(valor, list) else [valor]
            escribir(tf, pars, tam=tam, color=BLANCO if cab else NEGRO, negrita=cab)
        for j in range(len(fila) - 1, 0, -1):
            if fila[j] is None:
                k = j
                while k > 0 and fila[k] is None:
                    k -= 1
                t.cell(i, k).merge(t.cell(i, j))
                break
    return gf


# ---------------------------------------------------------- diapositivas ---
class Mazo:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Pt(960), Pt(540)
        self.vacia = self.prs.slide_layouts[6]

    def nueva(self, fondo, notas, numero=True, color_numero=BLANCO):
        s = self.prs.slides.add_slide(self.vacia)
        s.shapes.add_picture(str(REC / fondo), 0, 0, Pt(960), Pt(540))
        if numero:  # la portada no lleva número; la diapositiva del título es la 1
            self.numero(s, len(self.prs.slides) - 1, color_numero)
        s.notes_slide.notes_text_frame.text = notas
        return s

    @staticmethod
    def numero(s, n, color):
        caja = texto(s, 800, 452, 125, 72, str(n), tam=54, fuente='Calibri', color=color, alin='r')
        caja.name = 'Numero de pagina'
        if color == BLANCO:  # la sombra suave del número de la plantilla
            rPr = caja.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
            efecto = etree.Element(qn('a:effectLst'))
            sombra = etree.SubElement(efecto, qn('a:outerShdw'), blurRad='50800', dist='25400',
                                      dir='2700000', algn='tl', rotWithShape='0')
            clr = etree.SubElement(sombra, qn('a:prstClr'), val='black')
            etree.SubElement(clr, qn('a:alpha'), val='50000')
            rPr.find(qn('a:solidFill')).addnext(efecto)

    def contenido(self, titulo, notas, tam=40):
        s = self.nueva('fondo_contenido.jpg', notas)
        if len(titulo) * 0.6 * tam > 650:  # el título no debe llegar al logo
            tam = 650 / (len(titulo) * 0.6)
        texto(s, 89, 47, 660, 60, titulo, tam=tam, fuente='Arial Black', color=AZUL_TITULO)
        return s

    def separador(self, lineas, notas):
        s = self.nueva('fondo_separador.jpg', notas)
        texto(s, 96, 190, 660, 110, [{'t': l} for l in lineas], tam=40, fuente='Arial Black',
              color=BLANCO, interlineado=Pt(48))
        return s


# ----------------------------------------------------------- contenido ---
def portada(m):
    s = m.nueva('fondo_portada.png', 'Tiempo: 5 s.', numero=False)
    s.shapes.add_picture(str(REC / 'logo_portada.png'), Pt(119), Pt(122), Pt(279), Pt(271))
    texto(s, 126, 406, 700, 80, ['FACULTAD DE INGENIERÍA ELECTRÓNICA', 'INGENIERÍA ELECTRÓNICA'],
          tam=32, fuente='Calibri', color=BLANCO)


def titulo_proyecto(m):
    s = m.nueva('fondo_separador.jpg',
                'Tiempo: 20 s. Título del proyecto, autores y director. El trabajo pertenece al grupo '
                'GED, en la línea de Robótica e Inteligencia Computacional.')
    texto(s, 65, 64, 700, 140, ['SISTEMA COLABORATIVO DE ROBOTS', 'MÓVILES PARA ASISTENCIA DE',
                                'ORIENTACIÓN EN ENTORNOS INTERIORES', 'CON MÚLTIPLES PISOS'],
          tam=27, fuente='Arial Black', color=BLANCO, interlineado=Pt(34))
    texto(s, 68, 222, 400, 32, 'ESTUDIANTES:', tam=24, fuente='Arial Black', color=BLANCO)
    texto(s, 65, 265, 600, 64, ['Santiago Hernández Ávila', 'Jonny Alejandro Mejía León'], tam=24,
          negrita=True, color=BLANCO)
    texto(s, 68, 348, 700, 64, ['DIRECTOR: Ing. Armando Mateus Rojas, MSc.',
                                'LÍNEA: Robótica e Inteligencia Computacional'], tam=24,
          color=BLANCO)
    texto(s, 73, 479, 650, 32, 'Vinculado al grupo de investigación: GED', tam=24, negrita=True,
          color=BLANCO)


AGENDA = ['Introducción', 'Planteamiento del problema', 'Objetivos', 'Marco referencial',
          'Desarrollo', 'Resultados', 'Conclusiones', 'Trabajos futuros',
          'Referencias bibliográficas', 'Anexos']


def agenda(m):
    s = m.nueva('fondo_agenda.jpg',
                'Tiempo: 15 s. Orden de la presentación: el problema y los objetivos, el marco '
                'referencial, el desarrollo del sistema, los resultados de las dos campañas, las '
                'conclusiones y los trabajos futuros.')
    texto(s, 123, 46, 400, 55, 'CONTENIDO', tam=40, fuente='Arial Black', color=BLANCO)
    texto(s, 123, 96, 200, 24, 'Agenda', tam=18, negrita=True, color=AMARILLO)
    for i, item in enumerate(AGENDA):
        y = 122 + i * 38.6
        texto(s, 330, y, 378, 34, item, tam=24, color=BLANCO, alin='r', anclaje='m')
        forma(s, MSO_SHAPE.RECTANGLE, 733, y, 14, 33, relleno=AMARILLO)


def introduccion(m):
    m.separador(['INTRODUCCIÓN'], 'Tiempo: 3 s.')
    s = m.contenido('Introducción', (
        'Tiempo: 50 s. Orientarse en un edificio desconocido exige construir un mapa mental, y la '
        'escalera es el punto donde esa representación deja de servir (Jiménez Muñoz, 2024). Los '
        'trabajos que llevan un robot de un piso a otro lo hacen con hardware: mecanismos para subir '
        'escaleras o la integración con el ascensor (Kim et al., 2024; Song et al., 2024). Eso '
        'encarece la plataforma y obliga a intervenir el edificio. La propuesta usa dos agentes, '
        'uno por piso: el usuario cambia de piso por su cuenta y el agente del otro piso lo recibe. '
        'El cambio de piso pasa a ser un problema de coordinación, y ese traspaso, el relevo, es el '
        'aporte del proyecto. Se evaluó en simulación y en el edificio Catalayud.'))
    barras = ['Orientarse en un edificio desconocido es difícil',
              'La escalera corta la continuidad del recorrido',
              'Con un solo robot, cambiar de piso resulta caro',
              'Dos agentes, uno por piso, se relevan en la escalera']
    for i, t in enumerate(barras):
        y = 166 + i * 53.5
        forma(s, MSO_SHAPE.RECTANGLE, 64, y + 16, 529, 30, borde=AZUL)
        forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, 96, y, 447, 33, relleno=AZUL, radio=0.12,
              parrafos=t, tam=15, negrita=True, color=BLANCO, margen=(18, 0, 6, 0))
    texto(s, 64, 395, 500, 50, 'Plataforma AWS DeepRacer y ROS 2, evaluada en simulación con 30 '
          'misiones y en el edificio Catalayud con 10', tam=13, color=GRIS)
    s.shapes.add_picture(str(REC / 'imagen_introduccion.jpg'), Pt(593), Pt(98), Pt(221), Pt(319))
    fuente(s, 600, 418, 210, 'Fuente: Imagen generada con IA', alin='c')


def planteamiento(m):
    m.separador(['PLANTEAMIENTO DEL', 'PROBLEMA'], 'Tiempo: 3 s.')
    s = m.contenido('Planteamiento del problema', (
        'Tiempo: 50 s. Orientarse exige construir y actualizar una representación mental del '
        'espacio, y en un edificio de varios pisos el cambio de nivel rompe la continuidad del '
        'recorrido. La robótica colaborativa tiene una teoría de asignación de tareas aceptada desde '
        'Gerkey y Matarić (2004), pero la mayoría de sus desarrollos se evaluó en una sola planta '
        '(Nair, 2024). Resolver el cambio de piso con un solo robot exige subir escaleras o '
        'intervenir el ascensor, como hacen Song et al. (2024) con dispositivos conectados en la '
        'cabina: más costo de hardware y obras en el edificio. De ahí la pregunta problema.'))
    forma(s, MSO_SHAPE.DIAMOND, 102, 170, 320, 322, relleno=DURAZNO)
    cuadros = [(131, 200, 'Reto en la verticalidad'), (267, 200, 'Desafío de la navegación humana'),
               (131, 336, 'Barreras en sistemas con agente único'),
               (267, 336, 'Limitaciones económicas y tecnológicas')]
    for x, y, t in cuadros:
        forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 127, 127, relleno=BLANCO, borde=NARANJA_BORDE,
              radio=0.16, parrafos={'t': t, 'alin': 'c'}, tam=17, margen=5)
    forma(s, MSO_SHAPE.RIGHT_BRACE, 448, 170, 34, 338, borde=AZUL)
    forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, 520, 268, 322, 140, relleno=NARANJA, borde='F5B195',
          radio=0.18, parrafos='¿Cómo asistir la orientación de usuarios dentro de entornos '
          'interiores con múltiples pisos utilizando sistemas robóticos?', tam=20, negrita=True,
          margen=(14, 6, 10, 6))


def objetivos(m):
    m.separador(['OBJETIVOS'], 'Tiempo: 3 s.')
    s = m.contenido('Objetivo general', 'Tiempo: 20 s. Lectura del objetivo general, tal como lo '
                    'aprobó el anteproyecto.')
    forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, 107, 196, 703, 132, relleno=BLANCO, borde=AZUL_CLARO,
          radio=0.14, parrafos={'t': 'Implementar un sistema colaborativo de dos robots móviles '
                                'terrestres para la guía coordinada y eficiente de asistencia a '
                                'visitantes en entornos interiores con múltiples pisos, mediante '
                                'comunicación inter-robot, asignación dinámica de tareas y una '
                                'interfaz móvil de interacción usuario-robot.', 'alin': 'c'},
          tam=18, negrita=True, cursiva=True, margen=(12, 6, 12, 6))
    s.shapes.add_picture(str(REC / 'icono_objetivo.png'), Pt(400), Pt(352), Pt(120), Pt(124))

    s = m.contenido('Objetivos específicos', (
        'Tiempo: 40 s. Los cuatro objetivos específicos del anteproyecto, transcritos sin cambios: '
        'la arquitectura funcional, la plataforma con los dos agentes, la interfaz y la evaluación '
        'experimental con las cuatro métricas. El segundo nombra el DonkeyCar; la plataforma final '
        'fue el AWS DeepRacer, y el cambio se explica en el recorrido del trabajo.'))
    puntos = [(90 + 40 * math.sin(math.pi * (y - 128) / 412), y) for y in range(128, 541, 6)]
    linea(s, puntos, color=AZUL, grosor=1)
    items = [
        (89, 147, 182, 153, 'Modelar la arquitectura funcional del sistema colaborativo de dos robots '
         'móviles, definiendo requerimientos, escenarios de operación con múltiples pisos, '
         'estrategia de asignación de tareas y esquema de comunicación inter-robot.'),
        (127, 245, 219, 252, 'Desarrollar una plataforma robótica móvil basada en dos vehículos tipo '
         'DonkeyCar, integrando los módulos de locomoción, sensado, procesamiento y comunicación '
         'necesarios para la operación colaborativa en entornos interiores.'),
        (127, 345, 219, 352, 'Programar una interfaz móvil para interacción humano-robot (HRI) '
         'basada en la selección de localizaciones de interés de origen y destino.'),
        (89, 443, 182, 449, 'Evaluar el desempeño del sistema mediante pruebas experimentales en un '
         'entorno interior controlado, utilizando métricas como tiempo de respuesta, tiempo de '
         'asignación de robot, tasa de éxito en la entrega de asistencia y continuidad del servicio '
         'entre niveles.'),
    ]
    for k, (cx, cy, bx, by, t) in enumerate(items, 1):
        forma(s, MSO_SHAPE.RECTANGLE, bx - 12, by, 866 - bx + 12, 68 if k != 3 else 62,
              relleno=BLANCO, borde=AZUL)
        texto(s, bx, by + 3, 866 - bx - 6, 64, {'t': t, 'alin': 'j'}, tam=15.5, anclaje='m')
        forma(s, MSO_SHAPE.OVAL, cx, cy, 81, 81, relleno=BLANCO, borde=AZUL, grosor=1,
              parrafos={'t': str(k), 'alin': 'c'}, tam=30, fuente='Arial Black', color=AZUL_TITULO,
              margen=0)

    s = m.contenido('Alcance', (
        'Tiempo: 30 s. El alcance del sistema: una prueba de concepto con dos agentes, uno por piso, '
        'que se relevan en la escalera sin que ninguno cruce de nivel; una misión a la vez, en '
        'interiores y en condiciones controladas. Los agentes no accionan puertas ni ascensores y no '
        'transportan carga. La foto es de la campaña del 16 de octubre: una persona sigue al agente '
        'del piso 4.'))
    forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, 70, 128, 330, 372, relleno=BLANCO, borde=AZUL_ALCANCE,
          grosor=1.5, radio=0.08)
    forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, 110, 110, 250, 40, relleno=BLANCO, borde=AZUL_ALCANCE,
          grosor=1.5, radio=0.5, parrafos={'t': 'ALCANCE DEL SISTEMA', 'alin': 'c'}, tam=15,
          negrita=True)
    alcance = ['Desarrollo de una prueba de concepto funcional',
               'Implementación con dos agentes colaborativos, uno por piso',
               'Operación en entorno interior controlado, una misión a la vez',
               'Guía de usuarios entre pisos mediante relevo, sin que ningún agente cruce de piso',
               'Validación mediante pruebas experimentales en simulación y con los agentes físicos']
    for i, t in enumerate(alcance):
        y = 168 + i * 64
        texto(s, 92, y - 4, 26, 28, '✔', tam=20, fuente='Segoe UI Symbol', color=VERDE)
        texto(s, 118, y, 270, 56, t, tam=14)
    imagen(s, REC / 'foto_guiado.jpg', 430, 140, w=390)
    fuente(s, 430, 362, 390, 'Una persona sigue al agente del piso 4 en el edificio Catalayud, '
           '16 de octubre de 2026. Fuente: Elaboración propia')


def marco_referencial(m):
    m.separador(['MARCO REFERENCIAL'], 'Tiempo: 3 s.')
    s = m.contenido('Antecedentes', (
        'Tiempo: 50 s. De los robots guía se toma la integración de '
        'navegación e interacción; RHINO operó seis días en un museo en 1997. De la navegación en '
        'interiores, el flujo de Nav2 sobre ROS 2. De la coordinación multirrobot, la taxonomía de '
        'asignación y la planificación con relevos, donde lo que se traspasa es un paquete que puede '
        'esperar; una persona no puede, y por eso el relevo incluye una confirmación con plazo. De '
        'la movilidad entre pisos, el mapa por piso unido por un punto de transición; allí el robot '
        'cruza de nivel. La combinación nueva es resolver la discontinuidad entre pisos sin que '
        'ningún agente la atraviese.'))
    def linea_cita(nombre, cita):
        return [{'t': nombre, 'negrita': True, 'tam': 12}, {'t': cita, 'tam': 9.5, 'color': GRIS}]
    filas = [
        ['Línea', 'Lo que se toma de ella', 'Supuesto que no se comparte'],
        [linea_cita('Robots guía', 'Burgard et al. (1999); Guerreiro et al. (2019); Jiménez Muñoz '
                    '(2024)'), 'La integración de navegación e interacción; la localización de '
         'Monte Carlo', 'Un solo robot acompaña todo el recorrido'],
        [linea_cita('Navegación en interiores', 'Macenski et al. (2020)'), 'El flujo de mapa, '
         'localización y planificación sobre ROS 2 (Nav2)', 'Ninguno: se adopta, adaptado a la '
         'cinemática del agente'],
        [linea_cita('Coordinación multirrobot', 'Gerkey y Matarić (2004); Srivastava et al. (2025)'),
         'La taxonomía de asignación; la planificación con relevos',
         'Lo que se traspasa puede esperar sin límite'],
        [linea_cita('Movilidad entre pisos', 'Kim et al. (2024); Song et al. (2024)'),
         'El mapa por piso unido por un punto de transición', 'El robot cruza de nivel'],
        [linea_cita('Distribución del cómputo', 'Malavolta et al. (2023)'), 'La navegación a bordo '
         'y la coordinación en un nodo central', 'Ninguno'],
    ]
    tabla(s, 66, 112, [228, 258, 228], filas, alto=48, tam=12, altos={0: 26})
    forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, 66, 410, 640, 70, relleno=NARANJA, borde='F5B195',
          radio=0.2, parrafos='Aporte: ningún agente atraviesa la zona de transición. La continuidad '
          'entre pisos se obtiene con comunicación entre agentes, y no con mecanismos para escaleras '
          'ni con la intervención del ascensor.', tam=13.5, margen=(12, 4, 12, 4))

    s = m.contenido('Marco teórico', (
        'Tiempo: 50 s. En la taxonomía de Gerkey y Matarić, la asignación del sistema es ST–SR–IA: '
        'cada agente hace un tramo a la vez, cada tramo lo hace un solo agente y la decisión se toma '
        'al llegar la solicitud. El relevo acopla las agendas de los dos agentes, que es la clase XD '
        'de Korsah et al. ROS 2 se comunica sobre DDS, un estándar de comunicación de datos en tiempo '
        'real. Un vehículo tipo automóvil no gira sobre su eje; por eso se usan el A* híbrido con '
        'caminos de Reeds-Shepp y el controlador Regulated Pure Pursuit. La localización AMCL usa la '
        'odometría del láser, que en un pasillo uniforme no registra el avance. La evaluación usa el '
        'intervalo de Wilson, la falsabilidad de cada criterio y la norma ASTM F3244-21.'))
    forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, 131, 100, 284, 52, relleno=AZUL, radio=0.1,
          parrafos={'t': 'Fundamentos teóricos', 'alin': 'c'}, tam=24, fuente='Calibri',
          color=BLANCO)
    grupos = [
        ('Asignación de tareas y coordinación',
         ['Asignación ST–SR–IA (Gerkey y Matarić, 2004)',
          'Relevo: dependencias entre agendas (Korsah et al., 2013)',
          'Comunicación con ROS 2 (Macenski et al., 2022)']),
        ('Navegación de un vehículo tipo automóvil',
         ['A* híbrido y caminos de Reeds-Shepp (Dolgov et al., 2010)',
          'Regulated Pure Pursuit (Macenski et al., 2023)',
          'AMCL con odometría del láser (Jaimez et al., 2016)',
          'Pasillo uniforme: avance no observable (Zhang et al., 2016)']),
        ('Interacción y evaluación',
         ['rosbridge para el navegador (Crick et al., 2017)',
          'Intervalo de Wilson (1927) para la tasa de éxito',
          'Falsabilidad de cada criterio (Popper, 1959)',
          'Norma ASTM F3244-21 para la demostración física']),
    ]
    y, medios = 166, []
    for cab, vin in grupos:
        alto = 26 + 15.5 * len(vin)
        medios.append(y + alto / 2)
        linea(s, [(159, y + alto / 2), (188, y + alto / 2)], color='34599C')
        forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, 188, y, 378, alto, relleno=BLANCO, borde=AZUL,
              radio=0.1, parrafos=[{'t': cab, 'negrita': True, 'tam': 13}] +
              [{'t': v, 'vineta': True} for v in vin], tam=11.5, anclaje='t',
              margen=(7, 4, 4, 2))
        y += alto + 14
    linea(s, [(159, 152), (159, medios[-1])], color='34599C')
    imagen(s, REC / 'iconos_navegacion.jpg', 590, 104, w=220)
    imagen(s, REC / 'icono_hri.jpg', 612, 280, w=175)
    fuente(s, 590, 392, 220, 'Fuente: Imagen generada con IA', alin='c')

    s = m.contenido('Marco legal y ambiental', (
        'Tiempo: 20 s. El proyecto se enmarca en los CONPES 3975 y 4069, que piden a la academia '
        'prototipos con herramientas de la Cuarta Revolución Industrial. El alcance excluye que los '
        'agentes interactúen mecánicamente con el edificio. El impacto ambiental es mínimo: equipo '
        'institucional que ya existía y baterías que van a los canales de residuos electrónicos.'))
    tarjetas = [
        (70, 'Marco legal', [
            'CONPES 3975: Política Nacional para la Transformación Digital e Inteligencia Artificial',
            'CONPES 4069: Política Nacional de Ciencia, Tecnología e Innovación',
            'Los agentes no accionan puertas ni ascensores, no transportan carga y no cruzan la zona '
            'de transición']),
        (425, 'Marco ambiental', [
            'Equipo institucional que ya existía: el proyecto no compró los agentes',
            'Baterías pequeñas: una sesión de campo gasta una fracción de su carga',
            'Las baterías de polímero de litio van a los canales de disposición de residuos '
            'electrónicos']),
    ]
    for x, cab, vin in tarjetas:
        forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 125, 335, 290, relleno=BLANCO, borde=AZUL,
              radio=0.06)
        forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 125, 335, 44, relleno=AZUL, radio=0.2,
              parrafos={'t': cab, 'alin': 'c'}, tam=22, fuente='Calibri', color=BLANCO)
        texto(s, x + 16, 185, 305, 220, [{'t': v, 'vineta': True, 'despues': 10} for v in vin],
              tam=14)


def desarrollo(m):
    m.separador(['DESARROLLO'], 'Tiempo: 3 s.')
    s = m.contenido('Recorrido del trabajo', (
        'Tiempo: 45 s. El trabajo siguió las seis fases del anteproyecto entre febrero y octubre de '
        '2026. En abril se pasó del DonkeyCar al modelo del AWS DeepRacer, migrado a ROS 2 Humble. '
        'Entre agosto y septiembre se integró todo en simulación, con los pisos 2 y 3 del edificio '
        'Gregorio, y la campaña de treinta misiones corrió el 4 y el 5 de septiembre. Con los agentes '
        'físicos, en los pasillos uniformes de Gregorio la odometría no registraba el avance y el '
        'tránsito de personas dificultaba las pruebas. El 2 de octubre se pasó a los pisos 3 y 4 del '
        'edificio Catalayud, y la campaña física corrió el 16 de octubre.'))
    periodos = [
        ('Feb.–mar.', 'Planteamiento y requisitos', 'Anteproyecto y sustentación; requisitos '
         'funcionales y restricciones'),
        ('Abril', 'Diseño conceptual', 'Análisis del DonkeyCar; modelo del AWS DeepRacer en ROS 2 '
         'Humble'),
        ('May.–jun.', 'Desarrollo progresivo', 'Pasillo en Gazebo; seguidor de paredes, '
         'descartado; primer mapa con SLAM'),
        ('Agosto', 'Integración', 'Nav2 para dirección Ackermann; contrato de interfaces; primera '
         'misión con relevo'),
        ('Septiembre', 'Integración y validación', 'Campaña de 30 misiones en simulación; interfaz '
         'en el teléfono; segundo agente'),
        ('Octubre', 'Validación', 'Pisos 3 y 4 de Catalayud; unidad inercial; campaña física el 16 '
         'de octubre'),
    ]
    linea(s, [(60, 152), (790, 152)], color=AZUL, grosor=2)
    for i, (per, fase, hecho) in enumerate(periodos):
        cx = 115 + i * 121
        texto(s, cx - 58, 116, 116, 22, {'t': per, 'alin': 'c'}, tam=14, negrita=True,
              color=AZUL_TITULO)
        forma(s, MSO_SHAPE.OVAL, cx - 8, 144, 16, 16, relleno=AZUL, borde=BLANCO, grosor=1.5)
        texto(s, cx - 58, 166, 116, 34, {'t': fase, 'alin': 'c'}, tam=12, negrita=True, color=AZUL)
        texto(s, cx - 58, 202, 116, 100, {'t': hecho, 'alin': 'c'}, tam=11.5)
    fotos = [
        (REC / 'donkeycar.jpg', 62, 'DonkeyCar. Fuente: docs.donkeycar.com'),
        (REC / 'deepracer.jpg', 237, 'AWS DeepRacer. Fuente: aws.amazon.com/es/deepracer'),
        (FIG / 'simulacion_relevo_dos_procesos.png', 412, 'Misión con relevo en simulación. '
         'Fuente: Elaboración propia'),
        (REC / 'foto_escalera.jpg', 587, 'Campaña en Catalayud. Fuente: Elaboración propia'),
    ]
    for ruta, x, pie in fotos:
        bx, by, bw, bh = imagen(s, ruta, x, 318, w=160, h=112)
        fuente(s, x, by + bh + 3, 160, pie, alin='c')

    s = m.contenido('Arquitectura funcional', (
        'Tiempo: 55 s. Un coordinador central en el portátil decide y manda; dos agentes, uno por '
        'piso, ejecutan; la página web es la vía del usuario. Lo que cruza la escalera es la persona, '
        'y lo que se traspasa entre agentes es la responsabilidad de guiarla. Antes de escribir código '
        'se fijó un contrato de tres reglas. La tercera evita tener dos copias de la misma lógica en '
        'el coordinador y en la página. Los requisitos se reconstruyeron desde los objetivos: '
        'treinta y seis, cada uno con su prueba. El código del coordinador, del planificador y de la '
        'página es el mismo en simulación y con los agentes físicos; cambia la implementación de los '
        'bloques.'))
    bx, by, bw, bh = imagen(s, FIG / 'arquitectura_vehiculos.png', 64, 104, h=388)
    fuente(s, bx, by + bh + 2, bw, 'Arquitectura en la campaña con los agentes físicos. '
           'Fuente: Elaboración propia')
    texto(s, 572, 108, 255, 24, 'Contrato de interfaces', tam=16, negrita=True, color=AZUL_TITULO)
    reglas = ['El coordinador mueve a un agente solo con la acción de navegación de ROS 2, o con la '
              'media vuelta',
              'La interfaz habla solo con el coordinador, nunca con los agentes',
              'El texto que ve el usuario lo redacta el coordinador, y la interfaz lo muestra literal']
    texto(s, 572, 136, 255, 170, [{'t': f'{i}. {r}', 'despues': 7} for i, r in
                                  enumerate(reglas, 1)], tam=13)
    texto(s, 572, 304, 255, 24, 'Requisitos', tam=16, negrita=True, color=AZUL_TITULO)
    texto(s, 572, 330, 240, 70, '36 requisitos, cada uno con la prueba que lo declara cumplido: '
          '29 funcionales y 7 restricciones del alcance', tam=13)
    texto(s, 572, 398, 200, 60, 'Mismo código en simulación y con los agentes físicos', tam=13,
          negrita=True, color=AZUL)

    s = m.contenido('Planificación del relevo', (
        'Tiempo: 45 s. El relevo se planifica en una función pura, sin ROS 2: entra el catálogo de '
        'puntos y el par origen-destino, y salen los tramos. En el mismo piso son dos tramos del mismo '
        'agente; entre pisos, cuatro, con el relevo en el punto de transferencia de cada nivel. Así '
        'una prueba automática recorre las 930 solicitudes posibles del catálogo de simulación en '
        'milisegundos, y con ella se encontró un defecto que una corrida no habría mostrado. Con un '
        'agente por nivel, la asignación no necesita optimizador: el método húngaro devolvería la '
        'misma tabla. '
        'Abajo, las etapas que publica el coordinador y que sirven de marcas de tiempo.'))
    bx, by, bw, bh = imagen(s, FIG / 'planificador.png', 64, 104, w=445)
    fuente(s, bx, by + bh + 1, bw, 'Función que planifica una misión. Fuente: Elaboración propia')
    vin = ['Función pura, sin ROS 2: recibe el catálogo, el origen y el destino, y devuelve los '
           'tramos o un error con su motivo',
           'Mismo piso: dos tramos de un agente. Pisos distintos: cuatro tramos y un relevo',
           'Una prueba automática planifica las 930 solicitudes posibles (31 × 30 puntos) en '
           'milisegundos, sin simulador',
           'Un agente por nivel: la asignación es ST–SR–IA, y el método húngaro devolvería la '
           'misma tabla']
    texto(s, 522, 110, 318, 250, [{'t': v, 'vineta': True, 'despues': 8} for v in vin], tam=13.5)
    texto(s, 64, 350, 400, 18, 'Etapas de una misión, como las publica el coordinador', tam=11,
          color=GRIS)
    etapas = ['Recibida', 'Primer tramo', 'Transferencia', 'Esperando confirmación',
              'Segundo tramo', 'Completada']
    for i, e in enumerate(etapas):
        x = 64 + i * 112
        forma(s, MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON, x, 370, 120, 48,
              relleno=AZUL if i < 5 else AZUL_TITULO)
        # el texto va en una caja aparte: el área de texto de la flecha es más angosta que la forma
        texto(s, x + (20 if i else 6), 370, 86 if i else 92, 48, {'t': e, 'alin': 'c'}, tam=10.5,
              negrita=True, color=BLANCO, anclaje='m')
    texto(s, 64, 430, 640, 50, 'Con esa prueba se encontró un defecto: con el origen en el punto de '
          'transferencia de su nivel, el plan mandaba al mismo agente dos veces a la misma pose, y '
          'Nav2 devuelve éxito al instante ante un desplazamiento de cero metros.', tam=12,
          color=GRIS)

    s = m.contenido('Plataforma robótica', (
        'Tiempo: 55 s. Cada agente es un AWS DeepRacer con su tarjeta original, Ubuntu 24.04 y ROS 2 '
        'Jazzy, un LiDAR RPLIDAR A1M8 y la unidad inercial de la tarjeta. Las condiciones de la '
        'plataforma definieron el diseño. No gira sobre su eje, así que Nav2 se configuró con el '
        'planificador híbrido A* y Regulated Pure Pursuit, con marcha atrás. No avanza por debajo de '
        '0,40 m/s, lo que fija la tolerancia de llegada en 0,5 m. El rumbo del láser derivaba hasta '
        '5,9° por minuto, y con la unidad inercial bajó a 0,3° por minuto o menos. Y en pasillos de '
        '2,5 a 3,2 m la media vuelta de Nav2 abortó contra la pared, por lo que se escribió una '
        'maniobra propia. La red usa un punto de acceso de 5 GHz por piso.'))
    bx, by, bw, bh = imagen(s, REC / 'foto_agente.jpg', 66, 106, w=285)
    fuente(s, bx, by + bh + 2, bw, 'Agente del piso 4, 16 de octubre. Fuente: Elaboración propia')
    texto(s, 66, 337, 290, 22, 'Componentes de cada agente', tam=14, negrita=True,
          color=AZUL_TITULO)
    comp = ['AWS DeepRacer, escala 1:18, dirección Ackermann',
            'Intel Atom E3930, 4 GB; Ubuntu 24.04 y ROS 2 Jazzy',
            'LiDAR RPLIDAR A1M8: 360°, de 0,15 a 12 m',
            'Unidad inercial Bosch BMI160 de la tarjeta',
            'Coordinador en el portátil, en un contenedor']
    texto(s, 66, 361, 290, 120, [{'t': c, 'vineta': True, 'despues': 2} for c in comp], tam=12)
    tarjetas = [
        ('Navegación tipo automóvil', 'Planificador Smac (A* híbrido, Reeds-Shepp) y Regulated Pure '
         'Pursuit, con marcha atrás; sin la recuperación que gira sobre el eje'),
        ('Banda muerta de tracción', 'El agente no avanza por debajo de 0,40 m/s; la tolerancia de '
         'llegada se fijó en 0,5 m'),
        ('Rumbo con la unidad inercial', 'Filtro de Kalman extendido: un giro de 90° se midió como '
         '90,17° y 88,38°, con deriva de 0,3° por minuto o menos'),
        ('Media vuelta en varios tiempos', 'Maniobra propia, vigilada con el láser y el mapa; 24 '
         'medias vueltas en la campaña física'),
    ]
    for i, (cab, cuerpo) in enumerate(tarjetas):
        x, y = 372 + (i % 2) * 230, 106 + (i // 2) * 136
        forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 220, 126, relleno=BLANCO, borde=AZUL_CLARO,
              radio=0.1, parrafos=[{'t': cab, 'negrita': True, 'tam': 13.5, 'color': AZUL_TITULO,
                                    'despues': 4}, {'t': cuerpo}], tam=12, anclaje='t',
              margen=(9, 7, 7, 5))
    texto(s, 372, 384, 380, 60, 'Red: un punto de acceso de 5 GHz por piso; tiempo de ida y vuelta '
          'entre los agentes de 5,99 ms de mediana, frente al límite de 100 ms', tam=12.5,
          color=GRIS)

    s = m.contenido('Interfaz de interacción', (
        'Tiempo: 45 s. El objetivo pedía una interfaz móvil, y se construyó una página web que sirve '
        'el propio portátil: un visitante no va a instalar una aplicación para preguntar por un '
        'salón. Funciona sin salida a internet. El usuario elige el origen y el destino y sigue la '
        'misión con el texto que redacta el coordinador. La confirmación de cambio de piso y la '
        'cancelación se añadieron al ver el sistema en uso. Sin la confirmación, una misión podía '
        'darse por completada con la persona en el piso equivocado; la cancelación devuelve al agente '
        'a su escalera.'))
    x1, y1, w1, h1 = imagen(s, FIG / 'hri_seguimiento.png', 66, 106, h=330)
    x2, y2, w2, h2 = imagen(s, FIG / 'hri_confirmacion_piso.png', x1 + w1 + 12, 106, h=330)
    fuente(s, 66, 440, x2 + w2 - 66, 'Izquierda: seguimiento de una misión en simulación. Derecha: '
           'espera de la confirmación de cambio de piso. Fuente: Elaboración propia')
    vin = ['Página web servida por el portátil: cualquier teléfono, sin instalar nada y sin '
           'internet',
           'El usuario elige origen y destino, y sigue la misión con el texto del coordinador',
           'Confirmación de cambio de piso: aviso a los 60 s y misión fallida a los 120 s',
           'Cancelación: el agente vuelve a su escalera; en seis pruebas quedó a entre 0,07 y '
           '0,34 m de su punto']
    texto(s, x2 + w2 + 18, 112, 850 - (x2 + w2 + 18), 240,
          [{'t': v, 'vineta': True, 'despues': 9} for v in vin], tam=13)

    s = m.contenido('Protocolo experimental', (
        'Tiempo: 45 s. El protocolo se escribió antes de instrumentar el sistema, para que ni el '
        'instrumento ni los recorridos se eligieran después de ver los resultados. Cada métrica tiene '
        'un evento de inicio y uno de fin que registra una máquina. El éxito se mide contra la '
        'odometría del agente y no contra el resultado de Nav2, que en una prueba dio éxito con el '
        'agente a 0,296 m de la meta. Las misiones se sortearon con una semilla fija, y una corrida '
        'fallida cuenta como fallo salvo causas de una lista cerrada.'))
    filas = [
        ['Métrica', 'Inicio', 'Fin'],
        ['Tiempo de respuesta', 'El coordinador acepta la solicitud', 'Primer movimiento del agente '
         'asignado: tres muestras de odometría a 0,02 m/s o más'],
        ['Tiempo de asignación', 'Primera publicación de la etapa recibida', 'Primera publicación '
         'del primer tramo, con el agente asignado'],
        ['Tasa de éxito', 'Proporción de misiones exitosas sobre las corridas válidas', None],
        ['Continuidad entre niveles', 'El primer agente alcanza su punto de transferencia',
         'Primer movimiento del segundo agente'],
    ]
    tabla(s, 66, 108, [160, 255, 300], filas, alto=38, tam=12, altos={0: 24, 3: 28})
    cajas = [
        (66, 'Criterio de éxito', ['Pose final a 0,25 m o menos del destino en simulación, y a '
                                   '0,5 m o menos con los agentes físicos',
                                   'Medida con la odometría del agente, nunca con el resultado de '
                                   'Nav2',
                                   'Etapa completada sin pasar por fallida y, entre pisos, con '
                                   'relevo']),
        (420, 'Diseño de las campañas', ['Simulación: 30 misiones sorteadas, 15 en un piso y 15 '
                                         'entre pisos',
                                         'Agentes físicos: 10 misiones sorteadas entre los pisos 3 '
                                         'y 4',
                                         'Descartes solo por causas de una lista cerrada; si superan '
                                         'el 20 %, la campaña no vale']),
    ]
    for x, cab, vin in cajas:
        forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 300, 340, 150, relleno=BLANCO, borde=AZUL_CLARO,
              radio=0.08, parrafos=[{'t': cab, 'negrita': True, 'tam': 14, 'color': AZUL_TITULO,
                                     'despues': 4}] +
              [{'t': v, 'vineta': True, 'despues': 3} for v in vin], tam=12, anclaje='t',
              margen=(10, 7, 8, 4))
    texto(s, 66, 460, 640, 40, 'Instrumentación: grabación de ROS de cada misión, registro validado '
          'contra un esquema y un analizador que calcula las métricas con sus intervalos', tam=12,
          color=GRIS)


def grafico_odometria(s, x, y, w, h):
    """Fracción del avance real que registra la odometría, tabla de odometría del documento."""
    datos = [  # de abajo hacia arriba, como los dibuja el gráfico de barras
        ('Catalayud, piso 4 (14,57 m)', 103.3, True),
        ('Catalayud, piso 4 con la IMU (8,08 m)', 101.9, True),
        ('Catalayud, piso 4 (6,68 m)', 95.4, True),
        ('Gregorio, hall del piso 3 (6 m)', 73.0, False),
        ('Gregorio, pasillo simulado de 46,9 m', 37.9, False),
        ('Gregorio, pasillo simulado del piso 2', 5.7, False),
        ('Gregorio, pasillo simulado del piso 3', 1.3, False),
    ]
    vmax, x0, w0 = 125, 0.43, 0.54  # eje de valores y área de trazado, en fracción del gráfico
    # franja del criterio (error de 10 % o menos), detrás del gráfico
    bx = x + w * (x0 + w0 * 90 / vmax)
    bw = w * w0 * 20 / vmax
    forma(s, MSO_SHAPE.RECTANGLE, bx, y + 4, bw, h - 8, relleno='E2E2E2')
    texto(s, bx - 30, y - 16, bw + 60, 16, {'t': 'criterio: 90 a 110 %', 'alin': 'c'}, tam=10,
          color=GRIS)
    cd = CategoryChartData()
    cd.categories = [d[0] for d in datos]
    cd.add_series('Avance registrado', [d[1] for d in datos])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Pt(x), Pt(y), Pt(w), Pt(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.name = 'Arial'
    ch.font.size = Pt(11)
    plot = ch.plots[0]
    plot.gap_width = 55
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale = 0, vmax
    va.has_major_gridlines = False
    va.visible = False
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(11)
    ca.format.line.color.rgb = rgb('A6A6A6')
    serie = plot.series[0]
    for i, (_, v, cumple) in enumerate(datos):
        p = serie.points[i]
        p.format.fill.solid()
        p.format.fill.fore_color.rgb = rgb(AZUL if cumple else NARANJA_BORDE)
        dl = p.data_label
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        tf = dl.text_frame
        tf.text = f'{v:.1f} %'.replace('.', ',')
        f = tf.paragraphs[0].runs[0].font
        f.size, f.bold = Pt(11), True
        f.color.rgb = rgb(NEGRO)
    pa = ch._chartSpace.find('.//' + qn('c:plotArea'))
    lay = pa.find(qn('c:layout'))
    if lay is None:
        lay = etree.SubElement(pa, qn('c:layout'))
        pa.insert(0, lay)
    for e in list(lay):
        lay.remove(e)
    ml = etree.SubElement(lay, qn('c:manualLayout'))
    for tag, val in [('layoutTarget', 'inner'), ('xMode', 'edge'), ('yMode', 'edge'),
                     ('x', str(x0)), ('y', '0.02'), ('w', str(w0)), ('h', '0.96')]:
        etree.SubElement(ml, qn(f'c:{tag}')).set('val', val)
    for padre, despues_de in [(ch._chartSpace, ch._chartSpace.find(qn('c:chart'))), (pa, None)]:
        viejo = padre.find(qn('c:spPr'))
        if viejo is not None:
            padre.remove(viejo)
        sp = etree.Element(qn('c:spPr'))
        etree.SubElement(sp, qn('a:noFill'))
        etree.SubElement(etree.SubElement(sp, qn('a:ln')), qn('a:noFill'))
        despues_de.addnext(sp) if despues_de is not None else padre.append(sp)


def resultados(m):
    m.separador(['RESULTADOS'], 'Tiempo: 3 s.')
    s = m.contenido('Evaluación en simulación', (
        'Tiempo: 60 s. Treinta misiones sorteadas, quince en un piso y quince entre pisos, sin '
        'ningún descarte, de las que veintiséis fueron exitosas. Lo que se puede afirmar es que la tasa de '
        'éxito supera el 70 %, no que sea del 86,7 %: con treinta misiones el intervalo mide 24 '
        'puntos de ancho. Los intervalos de los dos grupos se solapan, así que no se comparan. Los '
        'cuatro fallos quedaron cortos, a entre 0,284 y 0,347 m del destino frente a 0,25 m: son de '
        'precisión de llegada, y en las treinta misiones el coordinador asignó el agente y ejecutó '
        'el relevo. El agente arranca en menos de medio segundo; la asignación tarda microsegundos.'))
    filas = [['Misiones', 'n', 'Éxitos', 'Tasa', 'Intervalo al 95 %'],
             ['Campaña completa', '30', '26', '86,7 %', '70,3 a 94,7 %'],
             ['Mismo piso, sin relevo', '15', '12', '80,0 %', '54,8 a 93,0 %'],
             ['Entre pisos, con relevo', '15', '14', '93,3 %', '70,2 a 98,8 %']]
    tabla(s, 66, 108, [175, 40, 58, 70, 125], filas, alto=26, tam=12.5)
    forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, 552, 108, 285, 104, relleno=NARANJA, borde='F5B195',
          radio=0.16, parrafos=[{'t': 'Con 95 % de confianza, la tasa de éxito supera el 70 %',
                                 'negrita': True, 'tam': 16, 'despues': 4},
                                {'t': 'Intervalo de Wilson; los dos grupos se solapan y no se '
                                      'comparan', 'tam': 11.5}], margen=(12, 4, 10, 4))
    bx, by, bw, bh = imagen(s, FIG / 'llegada_simulacion.png', 64, 222, w=575)
    fuente(s, bx, by + bh + 1, bw, 'Error de llegada de las 30 misiones; en naranja, las cuatro '
           'fuera de la tolerancia de 0,25 m. Fuente: Elaboración propia')
    texto(s, 650, 236, 150, 150, 'Los cuatro fallos quedaron cortos, a entre 0,284 y 0,347 m: '
          'precisión de llegada de la navegación, sin ningún fallo de coordinación', tam=12.5)
    tiles = [('0,1 a 0,4 s', 'Tiempo de respuesta, en pasos de 100 ms del reloj de la simulación'),
             ('< 100 ms', 'Tiempo de asignación; en banco de pruebas, mediana de 154 a 175 µs'),
             ('0,1 a 0,2 s', 'Intervalo del relevo en las 15 misiones entre pisos')]
    for i, (num, lab) in enumerate(tiles):
        forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, 64 + i * 212, 420, 200, 78, relleno=BLANCO,
              borde=AZUL_CLARO, radio=0.12,
              parrafos=[{'t': num, 'fuente': 'Arial Black', 'tam': 19, 'color': AZUL_TITULO,
                         'alin': 'c'}, {'t': lab, 'tam': 10.5, 'alin': 'c', 'color': GRIS}],
              margen=(6, 2, 6, 2))

    s = m.contenido('Odometría del agente', (
        'Tiempo: 40 s. El agente no tiene codificadores en las ruedas: su odometría sale de comparar '
        'barridos sucesivos del láser. Entre el 26 de agosto y el 2 de octubre no dio una medida '
        'utilizable: en un pasillo recto y uniforme cada barrido es casi igual al anterior y el '
        'avance no está en la señal. En Catalayud, con hornacinas, casilleros y quiebres, la misma '
        'odometría quedó dentro del 5 % sin cambiar un parámetro. El estimador es el mismo en todas '
        'las barras; cambió el entorno. Por eso se cambió de edificio.'))
    grafico_odometria(s, 64, 128, 550, 330)
    fuente(s, 64, 462, 550, 'Fracción del avance real, medido con flexómetro, que registra la '
           'odometría del láser. Fuente: Elaboración propia')
    vin = ['El mismo estimador en todas las barras: cambió el entorno',
           'En los pasillos uniformes de Gregorio el avance no está en la señal del láser',
           'En Catalayud, ocho cambios de perfil en 25 m: error dentro del 5 % sin cambiar un '
           'parámetro',
           'Criterio de la prueba: error de 10 % o menos en un recorrido de 5 m o más']
    texto(s, 628, 118, 205, 260, [{'t': v, 'vineta': True, 'despues': 8} for v in vin], tam=12.5)

    s = m.contenido('Precisión de llegada', (
        'Tiempo: 30 s. La prueba de navegación pide llegar de un punto a otro con Nav2 a 0,5 m o '
        'menos, medido con flexómetro. El 7 de octubre un agente recorrió tres tramos encadenados en '
        'el piso 4 y regresó a la salida sin que nadie lo tocara; las llegadas medidas quedaron entre '
        '0,06 y 0,45 m. El 8 de octubre el otro agente hizo cinco tramos en el piso 3, con dos medias '
        'vueltas, sin recuperaciones. Lo que limita la precisión es la banda muerta: el agente se '
        'aproxima a 0,40 m/s.'))
    bx, by, bw, bh = imagen(s, FIG / 'cadena_piso4.png', 64, 108, w=735)
    fuente(s, bx, by + bh + 1, bw, 'Cadena de tres tramos en el piso 4 y regreso a la salida, 7 de '
           'octubre. Fuente: Elaboración propia')
    filas = [['Tramo', 'Destino', 'Tiempo', 'Recuperaciones', 'A la meta', 'Medida'],
             ['1', 'Salón 403', '13,5 s', '1', '0,26 m', 'estimada'],
             ['2', 'Salón 402', '27,8 s', '1', '0,45 m', 'flexómetro'],
             ['3', 'Salón 401', '5,2 s', '0', '0,06 m', 'flexómetro'],
             ['Regreso', 'Salida', '62,8 s', '6', '0,34 m', 'flexómetro']]
    tabla(s, 64, 335, [62, 82, 62, 105, 72, 80], filas, alto=26, tam=12)
    texto(s, 540, 338, 200, 140, [{'t': 'Las llegadas medidas quedaron dentro de 0,5 m', 'vineta':
                                   True, 'despues': 8},
                                  {'t': 'Piso 3, 8 de octubre: cinco tramos con dos medias vueltas, '
                                        'sin recuperaciones', 'vineta': True}], tam=12.5)

    s = m.contenido('Campaña con los agentes físicos', (
        'Tiempo: 60 s. El 16 de octubre, en los pisos 3 y 4 de Catalayud, con un agente por piso y '
        'el coordinador en el portátil. Se sortearon diez misiones, cinco de subida y cinco de '
        'bajada, y se corrieron sin reubicar los agentes, como en la operación real. Las nueve que '
        'llegaron a correr se completaron, todas con relevo; la décima quedó interrumpida porque el '
        'portátil se quedó sin batería. La figura descompone cada misión en sus etapas; la franja '
        'rayada es la espera de la confirmación. Con nueve de nueve, la tasa de éxito supera el 70 %, '
        'la misma cota de la simulación. La llegada se midió con la localización de cada agente.'))
    tiles = [('9 de 9', 'misiones sorteadas completadas, todas con relevo'),
             ('87,6 s', 'tiempo total, mediana (52,3 a 105,4 s)'),
             ('0,87 s', 'tiempo de respuesta, mediana'),
             ('0,234 m', 'llegada según la localización, mediana (0,143 a 0,435 m)')]
    for i, (num, lab) in enumerate(tiles):
        forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, 66 + i * 182, 106, 172, 80, relleno=BLANCO,
              borde=AZUL_CLARO, radio=0.12,
              parrafos=[{'t': num, 'fuente': 'Arial Black', 'tam': 22, 'color': AZUL_TITULO,
                         'alin': 'c'}, {'t': lab, 'tam': 10.5, 'alin': 'c', 'color': GRIS}],
              margen=(6, 2, 6, 2))
    bx, by, bw, bh = imagen(s, FIG / 'campana_vehiculos_etapas.png', 64, 198, h=272)
    fuente(s, bx, by + bh + 1, bw, 'Duración de cada etapa en las nueve misiones. Fuente: '
           'Elaboración propia')
    vin = ['Intervalo de Wilson al 95 %: de 70,1 a 100 %',
           'ASTM F3244-21: nueve repeticiones sin fallo demuestran una fiabilidad mínima de 0,81 con '
           '85 % de confianza',
           'Tiempo de asignación de 0,3 ms; 24 medias vueltas de 160 a 171°',
           'La décima misión quedó interrumpida: el portátil se quedó sin batería']
    texto(s, bx + bw + 14, 204, 790 - (bx + bw + 14), 230,
          [{'t': v, 'vineta': True, 'despues': 7} for v in vin], tam=12)

    s = m.contenido('Relevo y limitaciones', (
        'Tiempo: 40 s. Para la persona, el relevo es el tiempo entre llegar a la escalera de su piso '
        'y que arranque el guiado en el otro: mediana de 24,8 s, en el que cambia de piso y confirma '
        'en la página. En cuanto a la medición, la continuidad, como criterio binario, no '
        'puede fallar por la forma en que está construida, y por eso se informa el intervalo del '
        'relevo. En el mapa de la simulación, celdas desconocidas se cargaron como libres; no cambia '
        'el veredicto y el umbral se corrigió el 25 de septiembre. Y con los agentes físicos la '
        'llegada se tomó de la localización del propio agente, sin flexómetro.'))
    bx, by, bw, bh = imagen(s, REC / 'foto_relevo.jpg', 66, 106, w=395)
    fuente(s, bx, by + bh + 2, bw, 'Misión 5: la persona baja del piso 4 y el agente del piso 3 la '
           'espera en su escalera. Fuente: Elaboración propia')
    forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, 66, 284, 395, 92, relleno=BLANCO, borde=AZUL_CLARO,
          radio=0.12, parrafos=[{'t': '24,8 s', 'fuente': 'Arial Black', 'tam': 24,
                                 'color': AZUL_TITULO, 'alin': 'c'},
                                {'t': 'mediana, de 18,0 a 29,5 s, desde que la persona llega a la '
                                      'escalera hasta que arranca el guiado en el otro piso',
                                 'tam': 11.5, 'alin': 'c', 'color': GRIS}], margen=(10, 2, 10, 2))
    texto(s, 66, 386, 395, 80, [{'t': 'La confirmación tomó de 2,3 a 24,7 s; el aviso de los 60 s '
                                      'no apareció en ninguna misión', 'vineta': True,
                                 'despues': 4},
                                {'t': 'Ninguna misión quedó sin agente asignado', 'vineta': True}],
          tam=12.5)
    texto(s, 478, 108, 330, 24, 'Limitaciones de la medición', tam=16, negrita=True,
          color=AZUL_TITULO)
    lim = ['La continuidad entre niveles, como criterio binario, no puede fallar por construcción: '
           'se informa el intervalo del relevo',
           'En el mapa de la simulación, 10 908 celdas desconocidas, el 14,5 % del mapa, se '
           'cargaron como libres; el umbral se corrigió el 25 de septiembre',
           'Con los agentes físicos, la llegada se midió con la localización de cada agente y no '
           'con flexómetro']
    texto(s, 478, 138, 325, 260, [{'t': f'{i}. {l}', 'despues': 10} for i, l in enumerate(lim, 1)],
          tam=13)


def conclusiones(m):
    m.separador(['CONCLUSIONES'], 'Tiempo: 3 s.')
    s = m.contenido('Conclusiones', (
        'Tiempo: 50 s. Los cuatro objetivos se cumplieron. Un agente por nivel convierte el cambio de '
        'piso en un problema de coordinación, y separar el planificador de ROS 2 permitió '
        'verificarlo completo. El AWS DeepRacer sirve como agente con tres condiciones que el trabajo '
        'hizo explícitas: el sitio de operación es parte del diseño, la tracción fija la precisión de '
        'llegada y el cómputo se reparte. Una página web basta para la interacción. Y la evaluación '
        'sostiene la cota del 70 % en las dos campañas, con los fallos en la navegación y no en la '
        'coordinación.'))
    concl = [
        'Un coordinador central con un agente por nivel convierte el cambio de piso en un problema '
        'de coordinación. La asignación es trivial por la estructura del problema; la dificultad '
        'está en coordinar dos agendas que dependen una de otra durante el relevo.',
        'El AWS DeepRacer sirve como agente con tres condiciones: su odometría depende del entorno, '
        'su tracción no responde por debajo de 0,40 m/s y su tarjeta obliga a coordinar desde un '
        'computador aparte.',
        'Una página web servida en la red local basta para pedir, seguir, confirmar y cancelar una '
        'misión desde cualquier teléfono, sin instalación y sin internet.',
        'Con 95 % de confianza, la tasa de éxito supera el 70 % en simulación, y la campaña física '
        'muestra la misma cota en un edificio real. Los fallos estuvieron en la precisión de '
        'llegada, no en la coordinación.',
    ]
    for i, t in enumerate(concl):
        y = 112 + i * 92
        forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, 118, y, 662 - (40 if i == 3 else 0), 82,
              relleno=BLANCO, borde=AZUL, radio=0.1,
              parrafos={'t': t, 'alin': 'j'}, tam=13.5, margen=(36, 4, 10, 4))
        forma(s, MSO_SHAPE.OVAL, 62, y + 4, 74, 74, relleno=BLANCO, borde=AZUL, grosor=1,
              parrafos={'t': str(i + 1), 'alin': 'c'}, tam=28, fuente='Arial Black',
              color=AZUL_TITULO, margen=0)

    s = m.contenido('Respuesta a la pregunta', (
        'Tiempo: 30 s. La respuesta a la pregunta problema es que un esquema de relevo entre dos '
        'agentes, uno por nivel, es viable. Los agentes se reparten el guiado de una persona que '
        'cambia de piso, sin que ninguno atraviese la zona de transición y sin modificar el '
        'edificio. Queda una conclusión de método: decidir qué mide cada criterio fue más difícil que '
        'construir el sistema. Falta medir la experiencia de la persona guiada, con un estudio con '
        'usuarios.'))
    forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, 66, 130, 306, 150, relleno=NARANJA, borde='F5B195',
          radio=0.16, parrafos='¿Cómo asistir la orientación de usuarios dentro de entornos '
          'interiores con múltiples pisos utilizando sistemas robóticos?', tam=18, negrita=True,
          margen=(14, 6, 12, 6))
    forma(s, MSO_SHAPE.RIGHT_ARROW, 382, 186, 56, 40, relleno=AZUL)
    forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, 450, 112, 360, 186, relleno=BLANCO, borde=AZUL_CLARO,
          grosor=1.25, radio=0.1,
          parrafos='Un esquema de relevo entre dos agentes, uno por nivel, es una respuesta viable: '
          'los agentes se reparten el guiado de una persona que cambia de piso sin que ninguno '
          'atraviese la zona de transición y sin modificar el edificio.', tam=15,
          margen=(14, 6, 12, 6))
    texto(s, 66, 318, 640, 140, [
        {'t': 'Decidir qué mide cada criterio fue más difícil que construir el sistema. Un criterio '
              'que no puede fallar, un reloj sin resolución para la magnitud que mide o un modelo del '
              'edificio sin verificar producen resultados que parecen evidencia y no lo son.',
         'vineta': True, 'despues': 10},
        {'t': 'Queda sin medir la experiencia de la persona guiada, que corresponde a un estudio con '
              'usuarios.', 'vineta': True}], tam=14)


def trabajos_futuros(m):
    m.separador(['TRABAJOS FUTUROS'], 'Tiempo: 3 s.')
    s = m.contenido('Trabajos futuros', (
        'Tiempo: 35 s. Una navegación más precisa, con tracción lenta y una medida del avance que no '
        'dependa del pasillo. Más capacidad de cómputo, con una Raspberry Pi 4 o 5. Evasión de '
        'obstáculos en movimiento, para operar en horario de clases. Un controlador predictivo como '
        'el MPPI de Nav2. E interacción por voz, con los modelos en el portátil, porque la red puede '
        'no tener internet.'))
    lineas = [
        ('Navegación más precisa', 'Calibrar la tracción o cambiar la transmisión, y medir el avance '
         'con codificadores u odometría visual, para volver a 0,25 m'),
        ('Más capacidad de cómputo', 'Llevar el agente a una Raspberry Pi 4 o 5, de cuatro núcleos; '
         'la comunidad del DeepRacer ya publica el software para la Pi 4'),
        ('Evasión de obstáculos', 'Esquivar personas y objetos en movimiento, y probarlo con gente '
         'en el pasillo, antes de operar en horario de clases'),
        ('Controlador predictivo', 'El MPPI de Nav2 admite cinemática tipo automóvil: la media '
         'vuelta podría volver a Nav2 y la llegada sería más precisa'),
        ('Interacción por voz', 'Voz a texto y un modelo de lenguaje en el portátil, que traduzcan la '
         'petición a un origen y un destino del catálogo'),
    ]
    for i, (cab, cuerpo) in enumerate(lineas):
        y = 116 + i * 70
        forma(s, MSO_SHAPE.RECTANGLE, 64, y + 16, 300, 30, borde=AZUL)
        forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, 88, y, 262, 33, relleno=AZUL, radio=0.12,
              parrafos=cab, tam=15, negrita=True, color=BLANCO, margen=(14, 0, 6, 0))
        texto(s, 380, y - 2, 410 - (30 if i == 4 else 0), 52, cuerpo, tam=13, anclaje='m')


REFERENCIAS = [
    ('ASTM International. (2021). ', 'Standard test method for navigation: Defined area', ' (ASTM '
     'F3244-21).'),
    ('Burgard, W., et al. (1999). Experiences with an interactive museum tour-guide robot. ',
     'Artificial Intelligence', '.'),
    ('Crick, C., et al. (2017). Rosbridge: ROS for non-ROS users. ', 'Robotics Research', '.'),
    ('Dolgov, D., Thrun, S., Montemerlo, M., y Diebel, J. (2010). Path planning for autonomous '
     'vehicles in unknown semi-structured environments. ', 'The International Journal of Robotics '
     'Research', '.'),
    ('Gerkey, B. P., y Matarić, M. J. (2004). A formal analysis and taxonomy of task allocation in '
     'multi-robot systems. ', 'The International Journal of Robotics Research', '.'),
    ('Guerreiro, J., et al. (2019). CaBot: Designing and evaluating an autonomous navigation robot '
     'for blind people. ', 'ASSETS', '.'),
    ('Jaimez, M., Monroy, J. G., y González-Jiménez, J. (2016). Planar odometry from a radial laser '
     'scanner. A range flow-based approach. ', 'ICRA', '.'),
    ('Jiménez Muñoz, M. (2024). Arquitectura cognitiva para integración de usuarios humanos dentro '
     'de un sistema de mapas multinivel.', '', ''),
    ('Kim, T., Kang, G., Lee, D., y Shim, D. H. (2024). Development of an indoor delivery mobile '
     'robot for a multi-floor environment. ', 'IEEE Access', '.'),
    ('Korsah, G. A., Stentz, A., y Dias, M. B. (2013). A comprehensive taxonomy for multi-robot task '
     'allocation. ', 'The International Journal of Robotics Research', '.'),
    ('Macenski, S., et al. (2020). The Marathon 2: A navigation system. ', 'IROS', '.'),
    ('Macenski, S., Foote, T., Gerkey, B., Lalancette, C., y Woodall, W. (2022). Robot Operating '
     'System 2: Design, architecture, and uses in the wild. ', 'Science Robotics', '.'),
    ('Macenski, S., et al. (2023). Regulated pure pursuit for robot path tracking. ',
     'Autonomous Robots', '.'),
    ('Malavolta, I., et al. (2023). Computation offloading for ground robotic systems communicating '
     'over WiFi. ', 'Empirical Software Engineering', '.'),
    ('Nair, B. R. (2024). Collaborative perception in multi-robot systems. ', 'ICRCV', '.'),
    ('Popper, K. R. (1959). ', 'The logic of scientific discovery', '. Hutchinson.'),
    ('Song, Z., Lou, S., Tan, R., y Lv, C. (2024). Robot-elevator interaction through IoT devices '
     'for autonomous robot multi-floor transition. ', 'EECR', '.'),
    ('Srivastava, A. K., Levin, J. M., y Dames, P. (2025). DELIVER: A system for coordinated '
     'multi-robot pickup and delivery using relay planning. ', 'IROS', '.'),
    ('Wilson, E. B. (1927). Probable inference, the law of succession, and statistical inference. ',
     'Journal of the American Statistical Association', '.'),
    ('Zhang, J., Kaess, M., y Singh, S. (2016). On degeneracy of optimization-based state '
     'estimation problems. ', 'ICRA', '.'),
]


def referencias(m):
    m.separador(['REFERENCIAS', 'BIBLIOGRÁFICAS'], 'Tiempo: 3 s.')
    s = m.contenido('Referencias bibliográficas', 'Tiempo: 5 s. Las referencias citadas en la '
                    'presentación; la lista completa, de 41 referencias, está en el documento.')
    mitad = (len(REFERENCIAS) + 1) // 2
    for col, refs in enumerate([REFERENCIAS[:mitad], REFERENCIAS[mitad:]]):
        pars = [{'runs': [(a, {}), (b, {'cursiva': True}), (c, {})], 'vineta': True, 'despues': 4}
                for a, b, c in refs]
        texto(s, 66 + col * 372, 108, 360 if col == 0 else 330, 400, pars, tam=10.5)


def anexos(m):
    m.separador(['ANEXOS'], 'Tiempo: 3 s.')
    s = m.contenido('Anexos', (
        'Tiempo: 15 s. Cada código abre un video del proyecto; los dos últimos son las campañas en '
        'simulación y con los agentes físicos. El repositorio guarda toda la trazabilidad: código, '
        'matriz de requisitos con la prueba de cada uno, protocolo, registro de cada misión, '
        'evidencia de cada sesión y entregables semanales.'))
    texto(s, 66, 104, 500, 22, 'Anexo A. Videos del proyecto', tam=15, negrita=True,
          color=AZUL_TITULO)
    for i, (vid, tema, fecha) in enumerate(VIDEOS):
        x = 66 + i * 102
        url = f'https://youtu.be/{vid}'
        imagen(s, REC / f'qr_video{i + 1}.png', x + 8, 130, w=80, h=80, enlace=url)
        texto(s, x, 213, 96, 60, [{'t': f'{i + 1}. {tema}', 'alin': 'c'},
                                  {'t': fecha, 'alin': 'c', 'color': GRIS}], tam=9.5)
    texto(s, 66, 290, 500, 22, 'Anexo B. Repositorio del proyecto', tam=15, negrita=True,
          color=AZUL_TITULO)
    imagen(s, REC / 'qr_repositorio.png', 66, 318, w=100, h=100, enlace=REPOSITORIO)
    texto(s, 176, 322, 220, 96, [{'t': 'Repositorio', 'negrita': True, 'tam': 13},
                                 {'t': 'github.com/shdeza12', 'color': GRIS},
                                 {'t': 'Código, requisitos con su prueba, protocolo, registro de '
                                       'cada misión, evidencia y entregables'}], tam=11)
    imagen(s, REC / 'qr_sitio.png', 410, 318, w=100, h=100, enlace=SITIO)
    texto(s, 520, 322, 200, 96, [{'t': 'Sitio del proyecto', 'negrita': True, 'tam': 13},
                                 {'t': 'shdeza12.github.io', 'color': GRIS},
                                 {'t': 'El mismo contenido, para leerlo en el navegador'}], tam=11)


def cierre(m):
    s = m.nueva('fondo_agenda.jpg', 'Tiempo: 5 s.', color_numero=NEGRO)
    forma(s, MSO_SHAPE.RECTANGLE, 180, 420, 765, 120, relleno=BLANCO)
    s.shapes.add_picture(str(REC / 'banda_acreditacion.png'), Pt(192), Pt(428), Pt(493), Pt(92))
    texto(s, 640, 334, 270, 56, {'t': 'GRACIAS', 'alin': 'r'}, tam=40, fuente='Arial Black',
          color=BLANCO)
    s = m.nueva('fondo_separador.jpg', 'El tiempo que quede es para las preguntas del jurado.')
    texto(s, 96, 190, 600, 70, '¿PREGUNTAS?', tam=48, fuente='Arial Black', color=BLANCO)


def main():
    m = Mazo()
    portada(m)
    titulo_proyecto(m)
    agenda(m)
    introduccion(m)
    planteamiento(m)
    objetivos(m)
    marco_referencial(m)
    desarrollo(m)
    resultados(m)
    conclusiones(m)
    trabajos_futuros(m)
    referencias(m)
    anexos(m)
    cierre(m)
    for s in m.prs.slides:  # el número de página, encima de todo, como en la plantilla
        for sh in list(s.shapes):
            if sh.name == 'Numero de pagina':
                s.shapes._spTree.remove(sh._element)
                s.shapes._spTree.append(sh._element)
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    m.prs.save(SALIDA)
    print(f'{SALIDA.relative_to(RAIZ)}: {len(m.prs.slides)} diapositivas')


if __name__ == '__main__':
    main()
